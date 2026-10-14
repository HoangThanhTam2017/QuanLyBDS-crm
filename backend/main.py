import io
import math
import os
from copy import copy
from pathlib import Path
from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent / ".env")

import base64
import hashlib
import re
import hmac
import json
import secrets
import urllib.error
import urllib.parse
import urllib.request
import uuid
from contextlib import asynccontextmanager
from datetime import date as date_type
from datetime import datetime, timedelta, timezone
from typing import List

import requests
import openpyxl
from openpyxl.styles import Font, PatternFill
from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from fastapi.responses import FileResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from jose import JWTError, jwt
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator
from sqlalchemy import case, func, inspect, or_, text
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session, joinedload

from database import Base, engine, get_db
from models import BusinessContract, Comment, LandPrice, PendingPost, Property, User

def load_local_env() -> None:
    backend_dir = Path(__file__).resolve().parent
    for env_path in (backend_dir.parent / ".env", backend_dir / ".env"):
        if not env_path.is_file():
            continue
        for line in env_path.read_text(encoding="utf-8").splitlines():
            item = line.strip()
            if not item or item.startswith("#") or "=" not in item:
                continue
            key, value = item.split("=", 1)
            key = key.strip()
            value = value.strip().strip('"').strip("'")
            if key and key not in os.environ:
                os.environ[key] = value


load_local_env()

JWT_SECRET = os.getenv("JWT_SECRET", "thanhtam-hatinh-dev-secret-change-me")
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
OPENAI_URL = "https://api.openai.com/v1/chat/completions"
COMPOSE_IRON_LAW = (
    "Bạn là một chuyên gia môi giới bất động sản lão luyện. "
    "Nhiệm vụ của bạn là viết lại bài đăng dựa trên thông tin được cung cấp. "
    "LUẬT THÉP: Tuyệt đối giữ nguyên vẹn và tính chính xác của 11 mục thông số kỹ thuật "
    "(Giá, Diện tích, Hướng, Vị trí...). Không dùng từ ngữ sáo rỗng (như 'bức tranh tuyệt đẹp'). "
    "Chỉ sáng tạo Tiêu đề, Mở bài (Hook) và Kết bài (Call-to-action). "
    "Giữ nguyên URL ở mục Link vị trí nếu có. "
    "Kết thúc bằng đúng hashtag: #BDSThanhTam #NhaDatHaTinh. "
    "Chỉ trả về bài viết, không markdown, không lời dẫn."
)
COMPOSE_STYLES = {
    "urgent": "Viết với nhịp điệu dồn dập, hối hả. Nhấn mạnh vào việc bán gấp, cắt lỗ, giá chạm đáy, cơ hội bắt đáy.",
    "invest": "Viết với văn phong sắc bén, phân tích. Nhấn mạnh vào dòng tiền, tỷ suất sinh lời, tiềm năng hạ tầng, vị trí đắc địa.",
    "living": "Viết với văn phong ấm áp, an tâm. Nhấn mạnh vào tiện ích xung quanh, an ninh, môi trường sống, tổ ấm gia đình.",
    "auto": "Tự động đọc Tiêu đề và Vị trí của bài viết để đoán định mục đích bán và áp dụng văn phong phù hợp nhất.",
}


def compose_system_prompt(style: str | None) -> str:
    key = (style or "auto").strip().lower()
    if key not in COMPOSE_STYLES:
        key = "auto"
    return (
        "Bạn là người viết tin đăng của Công ty cổ phần Thành Tâm Hà Tĩnh. "
        "Nhân viên gửi nội dung một bất động sản. Hãy viết lại thành bài đăng mạng xã hội.\n"
        + COMPOSE_IRON_LAW
        + "\n"
        + COMPOSE_STYLES[key]
    )
JWT_ALGORITHM = "HS256"
JWT_EXPIRE_HOURS = 8
STATIC_ROOT = Path(__file__).resolve().parent / "static"
FRONTEND_ROOT = Path(__file__).resolve().parent.parent / "frontend"
UPLOAD_DIR = STATIC_ROOT / "uploads"
ALLOWED_MEDIA = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".mp4", ".webm", ".mov"}
MAX_MEDIA_BYTES = 80 * 1024 * 1024
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def _column_is_numeric(column_type: object) -> bool:
    kind = str(column_type).lower()
    return any(token in kind for token in ("float", "double", "decimal", "numeric", "int", "real"))


def ensure_property_schema() -> None:
    """Bổ sung cột content/location và đổi price, area sang số nếu bảng cũ còn kiểu chữ."""
    inspector = inspect(engine)
    if "properties" not in inspector.get_table_names():
        return
    columns = {col["name"]: col for col in inspector.get_columns("properties")}
    statements: list[str] = []
    if "content" not in columns:
        statements.append("ALTER TABLE properties ADD COLUMN content TEXT NULL")
    if "location" not in columns:
        statements.append("ALTER TABLE properties ADD COLUMN location VARCHAR(255) NULL")
    if "media_urls" not in columns:
        statements.append("ALTER TABLE properties ADD COLUMN media_urls TEXT NULL")
    if "so_to" not in columns:
        statements.append("ALTER TABLE properties ADD COLUMN so_to VARCHAR(50) NULL")
    if "so_thua" not in columns:
        statements.append("ALTER TABLE properties ADD COLUMN so_thua VARCHAR(50) NULL")
    if "direction" not in columns:
        statements.append("ALTER TABLE properties ADD COLUMN direction VARCHAR(50) NULL")
    if "video_url" not in columns:
        statements.append("ALTER TABLE properties ADD COLUMN video_url VARCHAR(1000) NULL")
    if "description" in columns:
        statements.append(
            "UPDATE properties SET content = description "
            "WHERE (content IS NULL OR content = '') AND description IS NOT NULL AND description <> ''"
        )
    if "address" in columns:
        statements.append(
            "UPDATE properties SET location = LEFT(address, 255) "
            "WHERE (location IS NULL OR location = '') AND address IS NOT NULL AND address <> ''"
        )
    if "price" in columns and not _column_is_numeric(columns["price"]["type"]):
        statements.append(
            "UPDATE properties SET price = NULL "
            "WHERE price IS NOT NULL AND price NOT REGEXP '^[0-9]+([.][0-9]+)?$'"
        )
        statements.append("ALTER TABLE properties MODIFY COLUMN price DOUBLE NULL")
    if "area" in columns and not _column_is_numeric(columns["area"]["type"]):
        statements.append(
            "UPDATE properties SET area = NULL "
            "WHERE area IS NOT NULL AND area NOT REGEXP '^[0-9]+([.][0-9]+)?$'"
        )
        statements.append("ALTER TABLE properties MODIFY COLUMN area DOUBLE NULL")
    status_column = columns.get("status")
    if status_column is not None and "enum" in str(status_column["type"]).lower():
        statements.append("ALTER TABLE properties MODIFY COLUMN status VARCHAR(50) NULL DEFAULT 'Chưa bán'")
        statements.append(
            "UPDATE properties SET status = CASE status "
            "WHEN 'available' THEN 'Chưa bán' "
            "WHEN 'sold' THEN 'Đã bán' "
            "WHEN 'deposited' THEN 'Dừng bán' "
            "ELSE status END"
        )
    for name, spec in (
        ("title", "VARCHAR(255) NULL"),
        ("description", "TEXT NULL"),
        ("address", "VARCHAR(255) NULL"),
        ("direction", "VARCHAR(50) NULL"),
    ):
        column = columns.get(name)
        if column is not None and column.get("nullable") is False:
            statements.append(f"ALTER TABLE properties MODIFY COLUMN {name} {spec}")
    if not statements:
        return
    with engine.begin() as conn:
        for statement in statements:
            conn.execute(text(statement))


def ensure_pending_schema() -> None:
    """Thêm cột video_url cho bài chờ duyệt nếu bảng đã tồn tại từ trước."""
    inspector = inspect(engine)
    if "pending_posts" not in inspector.get_table_names():
        return
    columns = {col["name"] for col in inspector.get_columns("pending_posts")}
    if "video_url" in columns:
        return
    with engine.begin() as conn:
        conn.execute(text("ALTER TABLE pending_posts ADD COLUMN video_url VARCHAR(1000) NULL"))


def ensure_user_schema() -> None:
    """Bổ sung cột còn thiếu trên bảng users đã tạo từ trước."""
    inspector = inspect(engine)
    if "users" not in inspector.get_table_names():
        return
    columns = {col["name"] for col in inspector.get_columns("users")}
    statements: list[str] = []
    if "phone" not in columns:
        statements.append("ALTER TABLE users ADD COLUMN phone VARCHAR(20) NULL")
    if "trang_thai" not in columns:
        statements.append(
            "ALTER TABLE users ADD COLUMN trang_thai VARCHAR(50) NOT NULL DEFAULT 'Đang làm'"
        )
    else:
        statements.append(
            "UPDATE users SET trang_thai = 'Đang làm' WHERE trang_thai IS NULL OR trang_thai = ''"
        )
    if not statements:
        return
    with engine.begin() as conn:
        for statement in statements:
            conn.execute(text(statement))


@asynccontextmanager
async def lifespan(_app: FastAPI):
    try:
        Base.metadata.create_all(bind=engine)
        ensure_property_schema()
        ensure_pending_schema()
        ensure_user_schema()
    except SQLAlchemyError as exc:
        print(f"Chưa tạo được bảng users: {exc}")
    yield


app = FastAPI(title="Quản lý BĐS API", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "*",
        "http://localhost:5500",
        "http://127.0.0.1:5500",
    ],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

bearer_scheme = HTTPBearer(auto_error=False)


class LoginRequest(BaseModel):
    email: str = Field(min_length=1)
    password: str = Field(min_length=1)


class TokenResponse(BaseModel):
    access_token: str
    token_type: str = "bearer"
    id: int
    full_name: str = ""
    phone: str | None = None
    role: str = "user"
    nhom: str | None = None
    chuc_vu: str = "nhân viên"


class AccountRead(BaseModel):
    id: int
    email: str
    full_name: str
    phone: str | None = None
    avatar_url: str | None = None
    role: str = "user"
    nhom: str | None = None
    chuc_vu: str = "nhân viên"


class AccountUpdate(BaseModel):
    phone: str | None = None

    @field_validator("phone")
    @classmethod
    def clean_phone(cls, value: str | None) -> str | None:
        text = (value or "").strip()
        if not text:
            return None
        if len(text) > 20:
            raise ValueError("Số điện thoại quá dài.")
        return text


class ChangePasswordRequest(BaseModel):
    old_password: str
    new_password: str
    confirm_password: str


class MessageResponse(BaseModel):
    message: str


def account_read(user: User) -> AccountRead:
    return AccountRead(
        id=user.id,
        email=user.email,
        full_name=user.full_name or "",
        phone=user.phone,
        avatar_url=None,
        role=(user.role or "user").strip().lower(),
        nhom=user.nhom,
        chuc_vu=user.chuc_vu or "nhân viên",
    )


def create_access_token(user: User) -> str:
    expire = datetime.now(timezone.utc) + timedelta(hours=JWT_EXPIRE_HOURS)
    payload = {
        "sub": user.email,
        "name": user.full_name or "",
        "uid": user.id,
        "exp": expire,
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


ALLOWED_STATUS = ("Chưa bán", "Dừng bán", "Đã bán")
DIRECTION_PATTERN = re.compile(
    r"(Đông\s*Nam|Đông\s*Bắc|Tây\s*Nam|Tây\s*Bắc|Đông|Tây|Nam|Bắc)",
    re.IGNORECASE,
)
DIRECTION_NAMES = {
    "đông nam": "Đông Nam",
    "đông bắc": "Đông Bắc",
    "tây nam": "Tây Nam",
    "tây bắc": "Tây Bắc",
    "đông": "Đông",
    "tây": "Tây",
    "nam": "Nam",
    "bắc": "Bắc",
}


def normalize_direction(value: str | None) -> str | None:
    text = re.sub(r"\s+", " ", (value or "").strip())
    if not text:
        return None
    return DIRECTION_NAMES.get(text.lower())


def parse_direction(content: str | None) -> str | None:
    line = re.search(
        r"7\.\s*Phong\s*th(?:ủy|uỷ|uy)\s*[:\-]?\s*([^\n]*)",
        content or "",
        re.IGNORECASE,
    )
    if line is None:
        return None
    rest = line.group(1)
    pointed = re.search(r"hướng\s+", rest, re.IGNORECASE)
    match = DIRECTION_PATTERN.search(rest[pointed.end():]) if pointed else None
    if match is None:
        match = DIRECTION_PATTERN.search(rest)
    if match is None:
        return None
    return normalize_direction(match.group(1))


def clean_parcel(value: str | None) -> str | None:
    text = (value or "").strip()
    if not text:
        return None
    if len(text) > 50:
        raise ValueError("Số tờ hoặc số thửa quá dài.")
    return text


class PropertyCreate(BaseModel):
    content: str = Field(min_length=1)
    price: float | None = None
    area: float | None = None
    location: str = ""
    status: str = "Chưa bán"
    so_to: str | None = None
    so_thua: str | None = None
    direction: str | None = None

    @field_validator("status")
    @classmethod
    def default_status(cls, value: str) -> str:
        text = (value or "").strip()
        if not text:
            return "Chưa bán"
        if text not in ALLOWED_STATUS:
            raise ValueError("Trạng thái không hợp lệ.")
        return text

    @field_validator("so_to", "so_thua")
    @classmethod
    def parcel_value(cls, value: str | None) -> str | None:
        return clean_parcel(value)

    @field_validator("direction")
    @classmethod
    def direction_value(cls, value: str | None) -> str | None:
        return normalize_direction(value)


class ComposeRequest(BaseModel):
    content: str = Field(min_length=1, max_length=12000)
    style: str = "auto"

    @field_validator("style")
    @classmethod
    def style_value(cls, value: str) -> str:
        text = (value or "auto").strip().lower()
        if text not in COMPOSE_STYLES:
            return "auto"
        return text


class ComposeResponse(BaseModel):
    content: str


def strip_model_fences(text: str) -> str:
    cleaned = text.strip()
    if not cleaned.startswith("```"):
        return cleaned
    lines = cleaned.splitlines()
    if lines and lines[0].startswith("```"):
        lines = lines[1:]
    if lines and lines[-1].strip() == "```":
        lines = lines[:-1]
    return "\n".join(lines).strip()


def rewrite_with_openai(content: str, style: str | None = None) -> str:
    api_key = os.getenv("OPENAI_API_KEY", "").strip()
    if not api_key or api_key == "thay_api_key_cua_ban_vao_day":
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="Chưa cấu hình khóa OpenAI.")
    try:
        from openai import APIConnectionError, APIError, APITimeoutError, OpenAI
    except ImportError as exc:
        raise HTTPException(status_code=status.HTTP_503_SERVICE_UNAVAILABLE, detail="Chưa cài thư viện OpenAI.") from exc
    client = OpenAI(api_key=api_key, timeout=45)
    try:
        completion = client.chat.completions.create(
            model=OPENAI_MODEL or "gpt-4o-mini",
            temperature=0.7,
            messages=[
                {"role": "system", "content": compose_system_prompt(style)},
                {"role": "user", "content": content.strip()},
            ],
        )
        message = completion.choices[0].message.content if completion.choices else ""
        text = strip_model_fences(str(message or ""))
    except (APIError, APIConnectionError, APITimeoutError) as exc:
        print(f"OpenAI lỗi: {type(exc).__name__}")
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail="AI chưa viết được bài.") from exc
    except (AttributeError, IndexError, TypeError) as exc:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail="AI chưa viết được bài.") from exc
    if not text:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail="AI chưa viết được bài.")
    if "#BDSThanhTam" not in text:
        text += "\n\n#BDSThanhTam #NhaDatHaTinh"
    return text


class PropertyRead(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    content: str
    price: float | None
    area: float | None
    location: str
    so_to: str | None = None
    so_thua: str | None = None
    direction: str | None = None
    media_urls: list[str] = Field(default_factory=list)
    video_url: str | None = None
    status: str
    created_at: datetime
    user_id: int | None
    author_name: str


class DuplicateMatch(BaseModel):
    id: int
    title: str
    author_name: str


class CommentRead(BaseModel):
    id: int
    property_id: int
    user_id: int
    author_name: str
    content: str
    created_at: datetime


class CommentCreate(BaseModel):
    content: str

    @field_validator("content")
    @classmethod
    def clean_content(cls, value: str) -> str:
        text = (value or "").strip()
        if not text:
            raise ValueError("Nội dung bình luận không được để trống.")
        if len(text) > 2000:
            raise ValueError("Bình luận quá dài.")
        return text


def comment_read(item: Comment) -> CommentRead:
    author = item.author
    name = ""
    if author is not None:
        name = author.full_name or author.email
    return CommentRead(
        id=item.id,
        property_id=item.property_id,
        user_id=item.user_id,
        author_name=name or "Nhân viên",
        content=item.content,
        created_at=item.created_at,
    )


def read_media_urls(raw: object) -> list[str]:
    if raw is None:
        return []
    if isinstance(raw, list):
        return [str(item) for item in raw if item]
    text_value = str(raw).strip()
    if not text_value:
        return []
    try:
        parsed = json.loads(text_value)
    except json.JSONDecodeError:
        parsed = None
    if isinstance(parsed, list):
        return [str(item) for item in parsed if item]
    return [part.strip() for part in text_value.split(",") if part.strip()]


def optional_float(value: str, label: str) -> float | None:
    text_value = (value or "").strip()
    if not text_value:
        return None
    try:
        return float(text_value)
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"{label} không hợp lệ.",
        ) from exc


def save_media_files(uploads: List[UploadFile] | None) -> tuple[list[str], list[Path]]:
    urls: list[str] = []
    stored: list[Path] = []
    files = [item for item in (uploads or []) if item is not None and item.filename]
    if len(files) > 12:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mỗi tin tối đa 12 ảnh hoặc video.",
        )
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    try:
        for upload in files:
            suffix = Path(upload.filename or "").suffix.lower()
            if suffix not in ALLOWED_MEDIA:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Chỉ nhận ảnh (jpg, png, gif, webp) hoặc video (mp4, webm, mov).",
                )
            payload = upload.file.read()
            if len(payload) > MAX_MEDIA_BYTES:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Mỗi file tối đa 80MB.",
                )
            filename = f"{uuid.uuid4().hex}{suffix}"
            path = UPLOAD_DIR / filename
            path.write_bytes(payload)
            stored.append(path)
            urls.append(f"/static/uploads/{filename}")
    except (HTTPException, OSError):
        for path in stored:
            path.unlink(missing_ok=True)
        raise
    return urls, stored


VIDEO_SUFFIXES = {".mp4", ".webm", ".mov"}


def save_video_file(upload: UploadFile | None) -> tuple[str, list[Path]]:
    if upload is None or not (upload.filename or "").strip():
        return "", []
    suffix = Path(upload.filename or "").suffix.lower()
    if suffix not in VIDEO_SUFFIXES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Video chỉ nhận định dạng mp4, webm hoặc mov.",
        )
    urls, stored = save_media_files([upload])
    return (urls[0] if urls else ""), stored


IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".gif", ".webp"}


def classify_media_uploads(uploads: List[UploadFile] | None) -> tuple[list[UploadFile], UploadFile | None]:
    """Tách ảnh và video. Nhiều video thì chỉ giữ video đầu tiên."""
    images: list[UploadFile] = []
    video: UploadFile | None = None
    for upload in uploads or []:
        if upload is None or not (upload.filename or "").strip():
            continue
        content_type = (upload.content_type or "").split(";")[0].strip().lower()
        kind = ""
        if content_type.startswith("image/"):
            kind = "image"
        elif content_type.startswith("video/"):
            kind = "video"
        else:
            suffix = Path(upload.filename or "").suffix.lower()
            if suffix in IMAGE_SUFFIXES:
                kind = "image"
            elif suffix in VIDEO_SUFFIXES:
                kind = "video"
        if kind == "image":
            images.append(upload)
        elif kind == "video" and video is None:
            video = upload
    return images, video


def local_upload(url: str) -> Path | None:
    marker = "/static/uploads/"
    if not url.startswith(marker):
        return None
    name = url[len(marker):]
    if not name or "/" in name or "\\" in name or name.startswith("."):
        return None
    folder = UPLOAD_DIR.resolve()
    path = (folder / name).resolve()
    if path.parent != folder or not path.is_file():
        return None
    return path


PBKDF2_PREFIX = "pbkdf2_sha256"
PBKDF2_ITERATIONS = 260_000


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITERATIONS)
    return "$".join((
        PBKDF2_PREFIX,
        str(PBKDF2_ITERATIONS),
        base64.b64encode(salt).decode("ascii"),
        base64.b64encode(digest).decode("ascii"),
    ))


def password_matches(stored: str, provided: str) -> bool:
    if not stored or not provided:
        return False
    if stored.startswith(PBKDF2_PREFIX + "$"):
        try:
            _prefix, rounds, salt_b64, digest_b64 = stored.split("$", 3)
            salt = base64.b64decode(salt_b64)
            expected = base64.b64decode(digest_b64)
            actual = hashlib.pbkdf2_hmac("sha256", provided.encode("utf-8"), salt, int(rounds))
        except (ValueError, TypeError):
            return False
        if len(actual) != len(expected):
            return False
        return hmac.compare_digest(actual, expected)
    left = stored.encode("utf-8")
    right = provided.encode("utf-8")
    if len(left) != len(right):
        return False
    return hmac.compare_digest(left, right)


@app.post("/api/login", response_model=TokenResponse)
def login(body: LoginRequest, db: Session = Depends(get_db)):
    email = body.email.strip().lower()
    try:
        user = db.query(User).filter(func.lower(User.email) == email).first()
    except SQLAlchemyError as e:
        print(f"Chi tiết lỗi DB: {e}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from e

    if user is None or not password_matches(user.password_hash, body.password):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Email hoặc mật khẩu không đúng.",
            headers={"WWW-Authenticate": "Bearer"},
        )
    if (user.trang_thai or "").strip() == "Nghỉ việc":
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Tài khoản đã bị khóa, không thể đăng nhập",
        )

    return TokenResponse(
        access_token=create_access_token(user),
        id=user.id,
        full_name=user.full_name or "",
        phone=user.phone,
        role=(user.role or "user").strip().lower(),
        nhom=user.nhom,
        chuc_vu=user.chuc_vu or "nhân viên",
    )


def get_current_user(
    credentials: HTTPAuthorizationCredentials | None = Depends(bearer_scheme),
    db: Session = Depends(get_db),
) -> User:
    if credentials is None or credentials.scheme.lower() != "bearer":
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Thiếu token xác thực.",
            headers={"WWW-Authenticate": "Bearer"},
        )
    try:
        payload = jwt.decode(credentials.credentials, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        email = payload.get("sub")
    except JWTError as exc:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Token không hợp lệ hoặc đã hết hạn.",
            headers={"WWW-Authenticate": "Bearer"},
        ) from exc
    if not email:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Token không hợp lệ hoặc đã hết hạn.",
            headers={"WWW-Authenticate": "Bearer"},
        )
    try:
        user = db.query(User).filter(func.lower(User.email) == str(email).lower()).first()
    except SQLAlchemyError as e:
        print(f"Chi tiết lỗi DB: {e}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from e
    if user is None:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Token không hợp lệ hoặc đã hết hạn.",
            headers={"WWW-Authenticate": "Bearer"},
        )
    if (user.trang_thai or "").strip() == "Nghỉ việc":
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Tài khoản đã bị khóa, không thể đăng nhập",
            headers={"WWW-Authenticate": "Bearer"},
        )
    return user


class CurrentUserRead(BaseModel):
    full_name: str
    avatar_url: str | None = None


@app.get("/api/users/me", response_model=CurrentUserRead)
def read_users_me(current_user: User = Depends(get_current_user)):
    return CurrentUserRead(full_name=current_user.full_name or "", avatar_url=None)


@app.get("/api/me", response_model=AccountRead)
def read_me(current_user: User = Depends(get_current_user)):
    return account_read(current_user)


@app.patch("/api/me", response_model=AccountRead)
def update_me(
    body: AccountUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    current_user.phone = body.phone
    try:
        db.commit()
        db.refresh(current_user)
    except SQLAlchemyError as exc:
        db.rollback()
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không lưu được số điện thoại.",
        ) from exc
    return account_read(current_user)


@app.post("/api/users/change-password", response_model=MessageResponse)
def change_password(
    body: ChangePasswordRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if not password_matches(current_user.password_hash, body.old_password):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mật khẩu cũ không chính xác.",
        )
    if body.new_password != body.confirm_password:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mật khẩu mới và xác nhận mật khẩu không khớp.",
        )
    if not body.new_password.strip():
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mật khẩu mới không được để trống.",
        )
    current_user.password_hash = hash_password(body.new_password)
    try:
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không lưu được mật khẩu mới.",
        ) from exc
    return MessageResponse(message="Đổi mật khẩu thành công!")


class StaffOption(BaseModel):
    id: int
    full_name: str
    nhom: str | None = None


@app.get("/api/users/list", response_model=list[StaffOption])
def list_staff(
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    try:
        rows = db.query(User).order_by(User.full_name.asc(), User.id.asc()).all()
    except SQLAlchemyError as exc:
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from exc
    return [
        StaffOption(
            id=row.id,
            full_name=row.full_name or row.email or "Nhân viên",
            nhom=(row.nhom or "").strip() or None,
        )
        for row in rows
    ]


@app.get("/api/phuong-list")
def get_phuong_list(db: Session = Depends(get_db)):
    # Lấy danh sách phường không bị trùng lặp từ database
    phuongs = db.query(LandPrice.ma_phuong).distinct().order_by(LandPrice.ma_phuong.asc()).all()
    return [p[0] for p in phuongs if p[0]]


@app.get("/api/land-prices")
def search_land_prices(ma_phuong: str = None, tu_khoa: str = None, db: Session = Depends(get_db)):
    query = db.query(LandPrice)

    # Nếu có chọn phường cụ thể thì lọc theo phường
    if ma_phuong and ma_phuong != "ALL":
        query = query.filter(LandPrice.ma_phuong == ma_phuong)

    # Nếu có nhập tên đường thì dùng cú pháp LIKE không phân biệt hoa thường
    if tu_khoa:
        query = query.filter(LandPrice.duong.ilike(f"%{tu_khoa}%"))

    results = query.all()
    return results


LAM_TU_TYPES = ("Chính chủ", "CTV")
NGUON_TYPES = ("Nguồn cung", "Nguồn cầu")


class ContractRead(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    buyer_name: str
    seller_name: str
    land_address: str
    contract_quantity: int
    commission: float
    sale_price: float
    lam_tu: str
    nguon: str
    deposit_date: date_type
    ngay_thanh_toan: date_type | None
    status: str
    user_id: int | None = None
    ten_nhan_vien: str = ""


class ContractPage(BaseModel):
    data: list[ContractRead]
    total_pages: int
    current_page: int
    total: int
    user_role: str


class ReportHuy(BaseModel):
    so_luong: int
    ty_le: float


class ReportKpis(BaseModel):
    doanh_thu: float
    thanh_cong: int
    huy: ReportHuy
    nguon_khach: int


class ReportStaffRank(BaseModel):
    ten: str
    nhom: str
    so_hop_dong: int
    doanh_thu: float


class ReportGroupRank(BaseModel):
    ten_nhom: str
    so_hop_dong: int
    doanh_thu: float


class ReportCungCau(BaseModel):
    nhom: str
    nguon_cung: int
    nguon_cau: int


class ReportNguonSlice(BaseModel):
    nguon: str
    so_luong: int


class ReportLeaderboard(BaseModel):
    top_nhan_vien: list[ReportStaffRank]
    top_nhom: list[ReportGroupRank]


class ReportCharts(BaseModel):
    doanh_thu_thang: list[float]
    cung_cau_nhom: list[ReportCungCau]
    ty_le_nguon: list[ReportNguonSlice]


class ReportContractRow(BaseModel):
    id: int
    deposit_date: date_type
    ngay_thanh_toan: date_type | None
    land_address: str
    contract_quantity: int
    sale_price: float
    commission: float
    buyer_name: str
    seller_name: str
    lam_tu: str
    nguon: str
    status: str
    nhan_vien: str
    ten_nhan_vien: str
    nhom: str


class ReportDashboard(BaseModel):
    kpis: ReportKpis
    leaderboard: ReportLeaderboard
    charts: ReportCharts
    contracts: list[ReportContractRow]
    user_role: str


class ContractCreate(BaseModel):
    buyer_name: str
    seller_name: str
    land_address: str
    contract_quantity: int = Field(ge=1)
    commission: float = Field(ge=0)
    sale_price: float = Field(ge=0)
    lam_tu: str
    nguon: str
    deposit_date: date_type | None = None
    ngay_thanh_toan: date_type | None = None

    @field_validator("buyer_name", "seller_name", "land_address")
    @classmethod
    def clean_required_text(cls, value: str) -> str:
        text_value = value.strip()
        if not text_value:
            raise ValueError("Không được để trống.")
        return text_value

    @field_validator("lam_tu")
    @classmethod
    def clean_lam_tu(cls, value: str) -> str:
        text_value = value.strip()
        if text_value not in LAM_TU_TYPES:
            raise ValueError("Làm từ phải là Chính chủ hoặc CTV.")
        return text_value

    @field_validator("nguon")
    @classmethod
    def clean_nguon(cls, value: str) -> str:
        text_value = value.strip()
        if text_value not in NGUON_TYPES:
            raise ValueError("Nguồn phải là Nguồn cung hoặc Nguồn cầu.")
        return text_value


class ContractStatusUpdate(BaseModel):
    status: str

    @field_validator("status")
    @classmethod
    def clean_status(cls, value: str) -> str:
        text_value = value.strip()
        if text_value not in ("Đã thanh toán", "Hủy"):
            raise ValueError("Chỉ cập nhật sang Đã thanh toán hoặc Hủy.")
        return text_value


@app.get("/api/contracts", response_model=ContractPage)
def list_contracts(
    year: int = Query(),
    user_id: int | None = Query(default=None),
    status_name: str | None = Query(default=None, alias="status"),
    page: int = Query(default=1, ge=1),
    limit: int = Query(default=10, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    # user_id trên URL do frontend gửi kèm và không được dùng để nới quyền.
    del user_id
    query = db.query(BusinessContract).filter(func.extract("year", BusinessContract.deposit_date) == year)
    query = _contract_scope(query, current_user)
    if status_name is not None and status_name.strip():
        status_value = status_name.strip()
        if status_value not in ("Đặt cọc", "Đã thanh toán", "Hủy"):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Trạng thái không hợp lệ.")
        query = query.filter(BusinessContract.status == status_value)
    total_count = query.count()
    total_pages = math.ceil(total_count / limit) if total_count else 0
    contracts = (
        query.options(joinedload(BusinessContract.user))
        .order_by(BusinessContract.deposit_date.desc(), BusinessContract.id.desc())
        .offset((page - 1) * limit)
        .limit(limit)
        .all()
    )
    return ContractPage(
        data=[_contract_read(row) for row in contracts],
        total_pages=total_pages,
        current_page=page,
        total=total_count,
        user_role=_report_user_role(current_user),
    )


def _contract_scope(query, current_user: User):
    if _report_sees_company(current_user):
        return query
    return query.filter(BusinessContract.user_id == current_user.id)


def _contract_read(row: BusinessContract) -> ContractRead:
    item = ContractRead.model_validate(row)
    staff = row.user
    item.ten_nhan_vien = ((staff.full_name if staff else "") or "").strip()
    return item


def _report_title(current_user: User) -> str:
    return (current_user.chuc_vu or "").strip().lower()


def _report_sees_company(current_user: User) -> bool:
    role = (current_user.role or "").strip().lower()
    title = _report_title(current_user)
    return role == "admin" or title in ("admin", "giám đốc")


def _report_user_role(current_user: User) -> str:
    role = (current_user.role or "").strip().lower()
    title = _report_title(current_user)
    if role == "admin" or title == "admin":
        return "Admin"
    if title == "giám đốc":
        return "Giám đốc"
    return (current_user.chuc_vu or "").strip() or "nhân viên"


def _report_scope(
    query,
    current_user: User,
    view_mode: str = "personal",
    nhom: str | None = None,
    nhan_vien_id: int | None = None,
):
    if _report_sees_company(current_user):
        group_name = (nhom or "").strip()
        if group_name:
            query = query.filter(BusinessContract.nhom_phu_trach == group_name)
        if nhan_vien_id is not None:
            query = query.filter(BusinessContract.user_id == nhan_vien_id)
        return query
    mode = (view_mode or "personal").strip().lower()
    if mode == "group":
        group_name = (current_user.nhom or "").strip()
        if not group_name:
            return query.filter(BusinessContract.id == -1)
        return query.filter(BusinessContract.nhom_phu_trach == group_name)
    return query.filter(BusinessContract.user_id == current_user.id)


def _report_query(
    db: Session,
    year: int,
    month: int | None,
    current_user: User,
    view_mode: str = "personal",
    apply_month: bool = True,
    nhom: str | None = None,
    nhan_vien_id: int | None = None,
):
    query = db.query(BusinessContract).filter(func.extract("year", BusinessContract.deposit_date) == year)
    if apply_month and month is not None:
        query = query.filter(func.extract("month", BusinessContract.deposit_date) == month)
    return _report_scope(query, current_user, view_mode, nhom=nhom, nhan_vien_id=nhan_vien_id)


def _group_label(value: str | None) -> str:
    text = (value or "").strip()
    return text or "Chưa gán nhóm"


@app.get("/api/reports/dashboard", response_model=ReportDashboard)
def report_dashboard(
    year: int = Query(),
    month: int | None = Query(default=None, ge=1, le=12),
    view_mode: str = Query(default="personal"),
    nhom: str | None = Query(default=None),
    nhan_vien_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    scoped = _report_query(
        db, year, month, current_user, view_mode=view_mode, nhom=nhom, nhan_vien_id=nhan_vien_id
    )
    paid = scoped.filter(BusinessContract.status == "Đã thanh toán")
    doanh_thu = float(paid.with_entities(func.coalesce(func.sum(BusinessContract.commission), 0)).scalar() or 0)
    thanh_cong = int(paid.with_entities(func.count(BusinessContract.id)).scalar() or 0)
    huy_count = int(
        scoped.filter(BusinessContract.status == "Hủy").with_entities(func.count(BusinessContract.id)).scalar() or 0
    )
    total_count = int(scoped.with_entities(func.count(BusinessContract.id)).scalar() or 0)
    ty_le_huy = round(huy_count * 100 / total_count, 1) if total_count else 0
    nguon_khach = int(
        scoped.filter(func.lower(BusinessContract.lam_tu) == "ctv").with_entities(func.count(BusinessContract.id)).scalar() or 0
    )

    revenue = func.coalesce(func.sum(BusinessContract.commission), 0)
    staff_rows = (
        paid.join(User, BusinessContract.user_id == User.id)
        .with_entities(
            User.full_name,
            User.email,
            User.nhom,
            func.count(BusinessContract.id),
            revenue,
        )
        .group_by(User.id, User.full_name, User.email, User.nhom)
        .order_by(revenue.desc())
        .limit(5)
        .all()
    )
    top_nhan_vien = [
        ReportStaffRank(
            ten=(row[0] or row[1] or "Chưa có tên").strip(),
            nhom=_group_label(row[2]),
            so_hop_dong=int(row[3] or 0),
            doanh_thu=float(row[4] or 0),
        )
        for row in staff_rows
    ]

    group_rows = (
        paid.with_entities(
            BusinessContract.nhom_phu_trach,
            func.count(BusinessContract.id),
            revenue,
        )
        .group_by(BusinessContract.nhom_phu_trach)
        .order_by(revenue.desc())
        .all()
    )
    top_nhom = [
        ReportGroupRank(
            ten_nhom=_group_label(row[0]),
            so_hop_dong=int(row[1] or 0),
            doanh_thu=float(row[2] or 0),
        )
        for row in group_rows
    ]

    month_rows = (
        _report_query(
            db, year, month, current_user,
            view_mode=view_mode, apply_month=False, nhom=nhom, nhan_vien_id=nhan_vien_id,
        )
        .filter(BusinessContract.status == "Đã thanh toán")
        .with_entities(
            func.extract("month", BusinessContract.deposit_date),
            func.coalesce(func.sum(BusinessContract.commission), 0),
        )
        .group_by(func.extract("month", BusinessContract.deposit_date))
        .all()
    )
    doanh_thu_thang = [0.0] * 12
    for month_value, month_total in month_rows:
        index = int(month_value) - 1
        if 0 <= index < 12:
            doanh_thu_thang[index] = float(month_total or 0)

    direction_rows = (
        scoped.with_entities(
            BusinessContract.nhom_phu_trach,
            BusinessContract.nguon,
            func.count(BusinessContract.id),
        )
        .group_by(BusinessContract.nhom_phu_trach, BusinessContract.nguon)
        .all()
    )
    direction_map: dict[str, dict[str, int]] = {}
    for group_name, source_name, source_count in direction_rows:
        label = _group_label(group_name)
        bucket = direction_map.setdefault(label, {"nguon_cung": 0, "nguon_cau": 0})
        source = (source_name or "").strip().lower()
        if source == "nguồn cung":
            bucket["nguon_cung"] += int(source_count or 0)
        elif source == "nguồn cầu":
            bucket["nguon_cau"] += int(source_count or 0)
    cung_cau_nhom = [
        ReportCungCau(nhom=label, nguon_cung=values["nguon_cung"], nguon_cau=values["nguon_cau"])
        for label, values in sorted(direction_map.items(), key=lambda item: item[0].lower())
    ]

    source_rows = (
        scoped.with_entities(BusinessContract.lam_tu, func.count(BusinessContract.id))
        .group_by(BusinessContract.lam_tu)
        .order_by(func.count(BusinessContract.id).desc())
        .all()
    )
    ty_le_nguon = [
        ReportNguonSlice(nguon=(row[0] or "").strip() or "Khác", so_luong=int(row[1] or 0))
        for row in source_rows
    ]

    return ReportDashboard(
        kpis=ReportKpis(
            doanh_thu=doanh_thu,
            thanh_cong=thanh_cong,
            huy=ReportHuy(so_luong=huy_count, ty_le=ty_le_huy),
            nguon_khach=nguon_khach,
        ),
        leaderboard=ReportLeaderboard(top_nhan_vien=top_nhan_vien, top_nhom=top_nhom),
        charts=ReportCharts(
            doanh_thu_thang=doanh_thu_thang,
            cung_cau_nhom=cung_cau_nhom,
            ty_le_nguon=ty_le_nguon,
        ),
        contracts=[],
        user_role=_report_user_role(current_user),
    )


# Dòng 8 của template là tiêu đề cột. Tháng 1 bắt đầu ở dòng 9.
MONTH_ROW_MAP = {
    1: 9, 2: 10, 3: 11,
    4: 13, 5: 14, 6: 15,
    7: 17, 8: 18, 9: 19,
    10: 21, 11: 22, 12: 23,
}
QUARTER_ROW_MAP = {1: 12, 2: 16, 3: 20, 4: 24}
QUARTER_MONTHS = {1: (1, 2, 3), 2: (4, 5, 6), 3: (7, 8, 9), 4: (10, 11, 12)}
YEAR_TOTAL_ROW = 25
# Cung: D nguồn mới, F/G/H cọc, I/J/K GDTC, L nguồn CTV. Cầu: N nguồn mới, P/Q/R cọc, S/T/U GDTC, V nguồn CTV.
REPORT_VALUE_COLUMNS = (4, 6, 7, 8, 9, 10, 11, 12, 14, 16, 17, 18, 19, 20, 21, 22)


def _write_report_year(ws, year: int, col_offset: int = 0) -> None:
    cell = ws.cell(2, 9 + col_offset)
    title = cell.value
    if isinstance(title, str) and re.search(r"20\d{2}", title):
        cell.value = re.sub(r"20\d{2}", str(year), title, count=1)
    else:
        cell.value = "BẢNG KINH DOANH CÁ NHÂN NĂM " + str(year)


REPORT_STAT_FIELDS = (
    ("cung_sl", 4),
    ("cung_coc_sl", 6),
    ("cung_coc_gt", 7),
    ("cung_coc_hh", 8),
    ("cung_tt_sl", 9),
    ("cung_tt_gt", 10),
    ("cung_tt_hh", 11),
    ("cung_ctv", 12),
    ("cau_sl", 14),
    ("cau_coc_sl", 16),
    ("cau_coc_gt", 17),
    ("cau_coc_hh", 18),
    ("cau_tt_sl", 19),
    ("cau_tt_gt", 20),
    ("cau_tt_hh", 21),
    ("cau_ctv", 22),
)


def _empty_report_stats() -> dict[str, float]:
    return {key: 0 for key, _column in REPORT_STAT_FIELDS}


def _contract_month(contract: BusinessContract, report_year: int) -> int | None:
    ngay_dat_coc = contract.deposit_date
    if ngay_dat_coc is None:
        return None
    if isinstance(ngay_dat_coc, str):
        parts = str(ngay_dat_coc).split("-")
        nam = int(parts[0])
        thang = int(parts[1])
    else:
        nam = int(ngay_dat_coc.year)
        thang = int(ngay_dat_coc.month)
    if nam != report_year or thang not in MONTH_ROW_MAP:
        return None
    return thang


def _is_paid_ctv(contract: BusinessContract) -> bool:
    lam_tu = (contract.lam_tu or "").strip().lower()
    trang_thai = (contract.status or "").strip()
    return lam_tu == "ctv" and trang_thai == "Đã thanh toán"


def _accumulate_contract_stats(stats: dict[str, float], contract: BusinessContract) -> None:
    nguon = str(contract.nguon).lower().strip() if contract.nguon else ""
    trang_thai = str(contract.status).lower().strip() if contract.status else ""
    gia_ban = float(contract.sale_price or 0) / 1_000_000
    hoa_hong = float(contract.commission or 0) / 1_000_000
    if "cung" in nguon:
        stats["cung_sl"] += 1
        if "cọc" in trang_thai:
            stats["cung_coc_sl"] += 1
            stats["cung_coc_gt"] += gia_ban
            stats["cung_coc_hh"] += hoa_hong
        elif "thanh toán" in trang_thai or "thanh toan" in trang_thai:
            stats["cung_tt_sl"] += 1
            stats["cung_tt_gt"] += gia_ban
            stats["cung_tt_hh"] += hoa_hong
    elif "cầu" in nguon or "cau" in nguon:
        stats["cau_sl"] += 1
        if "cọc" in trang_thai:
            stats["cau_coc_sl"] += 1
            stats["cau_coc_gt"] += gia_ban
            stats["cau_coc_hh"] += hoa_hong
        elif "thanh toán" in trang_thai or "thanh toan" in trang_thai:
            stats["cau_tt_sl"] += 1
            stats["cau_tt_gt"] += gia_ban
            stats["cau_tt_hh"] += hoa_hong
    if _is_paid_ctv(contract):
        if "cung" in nguon:
            stats["cung_ctv"] += 1
        elif "cầu" in nguon or "cau" in nguon:
            stats["cau_ctv"] += 1


def _write_report_stats(ws, row: int, stats: dict[str, float], col_offset: int) -> None:
    for key, column in REPORT_STAT_FIELDS:
        value = stats[key]
        if value:
            ws.cell(row=row, column=column + col_offset).value = value
    quantity = stats["cung_tt_sl"] + stats["cau_tt_sl"]
    commission = stats["cung_tt_hh"] + stats["cau_tt_hh"]
    if quantity:
        ws.cell(row=row, column=23 + col_offset).value = quantity
        ws.cell(row=row, column=24 + col_offset).value = commission
        ws.cell(row=row, column=25 + col_offset).value = commission / quantity


def _sum_report_stats(items) -> dict[str, float]:
    total = _empty_report_stats()
    for stats in items:
        for key in total:
            total[key] += stats[key]
    return total


def _copy_row_style(ws, source_row: int, target_row: int, max_col: int = 27) -> None:
    for column in range(1, max_col + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        target.font = copy(source.font)
        target.border = copy(source.border)
        target.fill = copy(source.fill)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
        target.alignment = copy(source.alignment)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _write_staff_block(
    ws,
    row_idx: int,
    stats_by_user: dict,
    team_users: list[User],
    col_offset: int,
    tong_fill: PatternFill,
    bold_font: Font,
) -> None:
    staff_count = len(team_users)
    if staff_count:
        ws.insert_rows(row_idx + 1, staff_count)
        for offset in range(1, staff_count + 1):
            _copy_row_style(ws, row_idx, row_idx + offset)
    for index, user in enumerate(team_users):
        ws.cell(row=row_idx + index, column=3, value=(user.full_name or user.email or "").strip())
        _write_report_stats(ws, row_idx + index, stats_by_user[user.id], col_offset)
    total_row = row_idx + staff_count
    ws.cell(row=total_row, column=3, value="Tổng")
    _write_report_stats(
        ws,
        total_row,
        _sum_report_stats(stats_by_user[user.id] for user in team_users),
        col_offset,
    )
    for col_num in range(2, 27):
        cell = ws.cell(row=total_row, column=col_num)
        cell.fill = tong_fill
        cell.font = bold_font


def _write_group_month_rows(
    ws,
    data_month: dict,
    data_quarter: dict,
    team_users: list[User],
    tong_fill: PatternFill,
    bold_font: Font,
) -> None:
    # Dòng 8 là tiêu đề cột. Tháng 1 ở dòng 9, Quý IV ở dòng 24. Duyệt từ dưới lên.
    blocks = (
        ("Q", 4, 24), ("M", 12, 23), ("M", 11, 22), ("M", 10, 21),
        ("Q", 3, 20), ("M", 9, 19), ("M", 8, 18), ("M", 7, 17),
        ("Q", 2, 16), ("M", 6, 15), ("M", 5, 14), ("M", 4, 13),
        ("Q", 1, 12), ("M", 3, 11), ("M", 2, 10), ("M", 1, 9),
    )
    for kind, number, row_idx in blocks:
        source = data_quarter[number] if kind == "Q" else data_month[number]
        _write_staff_block(ws, row_idx, source, team_users, 1, tong_fill, bold_font)
    year_stats = _sum_report_stats(
        data_month[month][user.id] for month in range(1, 13) for user in team_users
    )
    _write_report_stats(ws, YEAR_TOTAL_ROW + len(blocks) * len(team_users), year_stats, 1)


def _fill_report_summaries(ws, col_offset: int = 0) -> None:
    columns = tuple(column + col_offset for column in REPORT_VALUE_COLUMNS)
    for quarter, months in QUARTER_MONTHS.items():
        target = QUARTER_ROW_MAP[quarter]
        for column in columns:
            total = sum((ws.cell(MONTH_ROW_MAP[month], column).value or 0) for month in months)
            if total:
                ws.cell(target, column).value = total
    for column in columns:
        total = sum((ws.cell(QUARTER_ROW_MAP[quarter], column).value or 0) for quarter in QUARTER_ROW_MAP)
        if total:
            ws.cell(YEAR_TOTAL_ROW, column).value = total
    for row_idx in list(MONTH_ROW_MAP.values()) + list(QUARTER_ROW_MAP.values()) + [YEAR_TOTAL_ROW]:
        quantity = (ws.cell(row_idx, 9 + col_offset).value or 0) + (ws.cell(row_idx, 19 + col_offset).value or 0)
        commission = (ws.cell(row_idx, 11 + col_offset).value or 0) + (ws.cell(row_idx, 21 + col_offset).value or 0)
        if quantity:
            ws.cell(row_idx, 23 + col_offset).value = quantity
            ws.cell(row_idx, 24 + col_offset).value = commission
            ws.cell(row_idx, 25 + col_offset).value = commission / quantity


@app.get("/api/contracts/export")
def export_contracts(
    year: int = Query(),
    export_type: str = Query(default="canhan"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    kind = (export_type or "canhan").strip().lower()
    if kind not in ("canhan", "nhom"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Loại xuất không hợp lệ.")
    if kind == "nhom":
        template_path = Path(__file__).resolve().parent / "BaoCaoNhom.xlsx"
        col_offset = 1
    else:
        template_path = Path(__file__).resolve().parent / "template.xlsx"
        col_offset = 0
    if not template_path.is_file():
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Không tìm thấy file mẫu báo cáo.")
    query = db.query(BusinessContract)
    if kind == "canhan":
        # Không phân trang. Admin và Giám đốc xuất mọi hợp đồng trong năm; các vai trò khác chỉ xuất hợp đồng của mình.
        if not _report_sees_company(current_user):
            query = query.filter(BusinessContract.user_id == current_user.id)
        scope_name = "Nhân viên: " + (current_user.full_name or current_user.email)
    else:
        chuc_vu = (current_user.chuc_vu or "").strip().lower()
        role = (current_user.role or "").strip().lower()
        if chuc_vu == "nhân viên" and role != "admin":
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Không có quyền")
        if not (current_user.nhom or "").strip():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Tài khoản chưa được gán nhóm.")
        query = query.filter(BusinessContract.nhom_phu_trach == current_user.nhom)
        scope_name = "Nhóm: " + current_user.nhom
    query = query.filter(func.extract("year", BusinessContract.deposit_date) == year)
    rows = query.all()
    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    tong_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    bold_font = Font(bold=True)
    ws["B2"] = scope_name
    _write_report_year(ws, year, col_offset)
    report_year = year
    if kind == "canhan":
        for contract in rows:
            thang = _contract_month(contract, report_year)
            if thang is None:
                continue
            row_idx = MONTH_ROW_MAP[thang]
            nguon = str(contract.nguon).lower().strip() if contract.nguon else ""
            trang_thai = str(contract.status).lower().strip() if contract.status else ""
            gia_ban = float(contract.sale_price or 0) / 1_000_000
            hoa_hong = float(contract.commission or 0) / 1_000_000
            if "cung" in nguon:
                ws.cell(row=row_idx, column=4 + col_offset).value = (ws.cell(row=row_idx, column=4 + col_offset).value or 0) + 1
                if "cọc" in trang_thai:
                    ws.cell(row=row_idx, column=6 + col_offset).value = (ws.cell(row=row_idx, column=6 + col_offset).value or 0) + 1
                    ws.cell(row=row_idx, column=7 + col_offset).value = (ws.cell(row=row_idx, column=7 + col_offset).value or 0) + gia_ban
                    ws.cell(row=row_idx, column=8 + col_offset).value = (ws.cell(row=row_idx, column=8 + col_offset).value or 0) + hoa_hong
                elif "thanh toán" in trang_thai or "thanh toan" in trang_thai:
                    ws.cell(row=row_idx, column=9 + col_offset).value = (ws.cell(row=row_idx, column=9 + col_offset).value or 0) + 1
                    ws.cell(row=row_idx, column=10 + col_offset).value = (ws.cell(row=row_idx, column=10 + col_offset).value or 0) + gia_ban
                    ws.cell(row=row_idx, column=11 + col_offset).value = (ws.cell(row=row_idx, column=11 + col_offset).value or 0) + hoa_hong
            elif "cầu" in nguon or "cau" in nguon:
                ws.cell(row=row_idx, column=14 + col_offset).value = (ws.cell(row=row_idx, column=14 + col_offset).value or 0) + 1
                if "cọc" in trang_thai:
                    ws.cell(row=row_idx, column=16 + col_offset).value = (ws.cell(row=row_idx, column=16 + col_offset).value or 0) + 1
                    ws.cell(row=row_idx, column=17 + col_offset).value = (ws.cell(row=row_idx, column=17 + col_offset).value or 0) + gia_ban
                    ws.cell(row=row_idx, column=18 + col_offset).value = (ws.cell(row=row_idx, column=18 + col_offset).value or 0) + hoa_hong
                elif "thanh toán" in trang_thai or "thanh toan" in trang_thai:
                    ws.cell(row=row_idx, column=19 + col_offset).value = (ws.cell(row=row_idx, column=19 + col_offset).value or 0) + 1
                    ws.cell(row=row_idx, column=20 + col_offset).value = (ws.cell(row=row_idx, column=20 + col_offset).value or 0) + gia_ban
                    ws.cell(row=row_idx, column=21 + col_offset).value = (ws.cell(row=row_idx, column=21 + col_offset).value or 0) + hoa_hong
            if _is_paid_ctv(contract):
                ctv_column = 12 if "cung" in nguon else 22 if ("cầu" in nguon or "cau" in nguon) else None
                if ctv_column is not None:
                    cell = ws.cell(row=row_idx, column=ctv_column + col_offset)
                    cell.value = (cell.value or 0) + 1
        _fill_report_summaries(ws, col_offset)
    else:
        current_team = db.query(User).filter(User.nhom == current_user.nhom).all()
        past_id_rows = (
            db.query(BusinessContract.user_id)
            .filter(
                BusinessContract.nhom_phu_trach == current_user.nhom,
                func.extract("year", BusinessContract.deposit_date) == year,
                BusinessContract.user_id.isnot(None),
            )
            .distinct()
            .all()
        )
        past_ids = [row[0] for row in past_id_rows]
        past_team = db.query(User).filter(User.id.in_(past_ids)).all() if past_ids else []
        team_users = list({user.id: user for user in current_team + past_team}.values())
        team_users.sort(key=lambda user: ((user.full_name or "").lower(), user.id))
        data_month = {month: {user.id: _empty_report_stats() for user in team_users} for month in range(1, 13)}
        data_quarter = {quarter: {user.id: _empty_report_stats() for user in team_users} for quarter in range(1, 5)}
        for contract in rows:
            thang = _contract_month(contract, report_year)
            if thang is None or contract.user_id not in data_month[thang]:
                continue
            _accumulate_contract_stats(data_month[thang][contract.user_id], contract)
            quy = (thang - 1) // 3 + 1
            _accumulate_contract_stats(data_quarter[quy][contract.user_id], contract)
        _write_group_month_rows(ws, data_month, data_quarter, team_users, tong_fill, bold_font)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = "BaoCaoNhom.xlsx" if kind == "nhom" else "BaoCao_KinhDoanh.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def reject_foreign_contract_edit(current_user: User, row: BusinessContract) -> None:
    if is_system_admin(current_user):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Admin chỉ có quyền xem, không được phép thao tác trên hợp đồng",
        )
    if row.user_id is not None and row.user_id == current_user.id:
        return
    raise HTTPException(
        status_code=status.HTTP_403_FORBIDDEN,
        detail="Bạn không có quyền thao tác trên hợp đồng của người khác",
    )


@app.post("/api/contracts", response_model=ContractRead, status_code=status.HTTP_201_CREATED)
def create_contract(
    body: ContractCreate,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    row = BusinessContract(
        buyer_name=body.buyer_name,
        seller_name=body.seller_name,
        land_address=body.land_address,
        contract_quantity=body.contract_quantity,
        commission=body.commission,
        sale_price=body.sale_price,
        lam_tu=body.lam_tu,
        nguon=body.nguon,
        deposit_date=body.deposit_date or date_type.today(),
        ngay_thanh_toan=body.ngay_thanh_toan,
        status="Đặt cọc",
        user_id=_user.id,
        nhom_phu_trach=_user.nhom,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


@app.put("/api/contracts/{contract_id}", response_model=ContractRead)
def update_contract(
    contract_id: int,
    body: ContractCreate,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    row = db.get(BusinessContract, contract_id)
    if row is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy hợp đồng.")
    reject_foreign_contract_edit(_user, row)
    row.buyer_name = body.buyer_name
    row.seller_name = body.seller_name
    row.land_address = body.land_address
    row.contract_quantity = body.contract_quantity
    row.commission = body.commission
    row.sale_price = body.sale_price
    row.lam_tu = body.lam_tu
    row.nguon = body.nguon
    row.deposit_date = body.deposit_date or row.deposit_date
    row.ngay_thanh_toan = body.ngay_thanh_toan
    db.commit()
    db.refresh(row)
    return row


@app.delete("/api/contracts/{contract_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_contract(
    contract_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    row = db.get(BusinessContract, contract_id)
    if row is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy hợp đồng.")
    reject_foreign_contract_edit(_user, row)
    db.delete(row)
    db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@app.put("/api/contracts/{contract_id}/status", response_model=ContractRead)
def update_contract_status(
    contract_id: int,
    body: ContractStatusUpdate,
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    row = db.get(BusinessContract, contract_id)
    if row is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy hợp đồng.")
    reject_foreign_contract_edit(_user, row)
    if row.status != "Đặt cọc":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Chỉ hợp đồng đang đặt cọc mới đổi được trạng thái.")
    row.status = body.status
    db.commit()
    db.refresh(row)
    return row


DEFAULT_STAFF_PASSWORD = "ThanhTam@123"


class StaffAccountRead(BaseModel):
    id: int
    full_name: str
    email: str
    role: str
    nhom: str | None = None
    chuc_vu: str = "nhân viên"
    trang_thai: str = "Đang làm"


class StaffAccountCreate(BaseModel):
    full_name: str = Field(min_length=1, max_length=255)
    email: str = Field(min_length=3, max_length=255)
    password: str = Field(min_length=6, max_length=128)
    role: str = "user"
    nhom: str | None = None
    chuc_vu: str = "nhân viên"
    trang_thai: str = "Đang làm"


class StaffAccountUpdate(BaseModel):
    full_name: str = Field(min_length=1, max_length=255)
    email: str = Field(min_length=3, max_length=255)
    role: str = "user"
    nhom: str | None = None
    chuc_vu: str = "nhân viên"
    trang_thai: str = "Đang làm"


class ResetPasswordResponse(BaseModel):
    message: str
    default_password: str


def is_system_admin(user: User) -> bool:
    role = (user.role or "").strip().lower()
    title = (user.chuc_vu or "").strip().lower()
    return role == "admin" or title == "admin"


def require_module_admin(current_user: User = Depends(get_current_user)) -> User:
    if not is_system_admin(current_user):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Bạn không có quyền truy cập",
        )
    return current_user


def require_staff_admin(current_user: User = Depends(get_current_user)) -> User:
    return require_module_admin(current_user)


def clean_staff_email(value: str) -> str:
    email = value.strip().lower()
    if "@" not in email or email.startswith("@") or email.endswith("@") or " " in email:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Email không hợp lệ.")
    return email


STAFF_GROUPS = ("Bạch Dương", "Thần Tài", "Mai Vàng")


def clean_nhom(value: str | None, db: Session | None = None) -> str | None:
    if value is None:
        return None
    group = value.strip()
    if not group:
        return None
    if group in STAFF_GROUPS:
        return group
    if db is not None:
        match = (
            db.query(User.nhom)
            .filter(User.nhom.isnot(None), User.nhom != "", func.lower(User.nhom) == group.lower())
            .first()
        )
        stored = (match[0] or "").strip() if match else ""
        if stored:
            return stored
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Nhóm không hợp lệ.")


def clean_staff_role(value: str) -> str:
    role = value.strip().lower()
    if role == "admin":
        return "admin"
    if role == "user":
        return "user"
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Quyền chỉ được là Admin hoặc User.")


STAFF_TITLES = ("nhân viên", "trưởng nhóm")


STAFF_STATUSES = ("Đang làm", "Nghỉ việc")


def clean_trang_thai(value: str | None) -> str:
    raw = (value or "").strip()
    if not raw:
        return "Đang làm"
    if raw in STAFF_STATUSES:
        return raw
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Trạng thái không hợp lệ.")


def clean_chuc_vu(value: str | None, db: Session | None = None) -> str:
    raw = (value or "").strip()
    if not raw:
        return "nhân viên"
    lowered = raw.lower()
    if lowered in STAFF_TITLES:
        return lowered
    if db is not None:
        match = (
            db.query(User.chuc_vu)
            .filter(User.chuc_vu.isnot(None), func.lower(User.chuc_vu) == lowered)
            .first()
        )
        stored = (match[0] or "").strip() if match else ""
        if stored:
            return stored
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Chức vụ không hợp lệ.")


def staff_account_read(user: User) -> StaffAccountRead:
    role = "Admin" if (user.role or "").strip().lower() == "admin" else "User"
    return StaffAccountRead(
        id=user.id,
        full_name=user.full_name or "",
        email=user.email,
        role=role,
        nhom=user.nhom or None,
        chuc_vu=user.chuc_vu or "nhân viên",
        trang_thai=user.trang_thai or "Đang làm",
    )


def find_staff_user(db: Session, user_id: int) -> User:
    user = db.get(User, user_id)
    if user is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy nhân viên.")
    return user


def admin_count(db: Session, exclude_id: int | None = None) -> int:
    query = db.query(User).filter(func.lower(User.role) == "admin")
    if exclude_id is not None:
        query = query.filter(User.id != exclude_id)
    return query.count()


@app.get("/api/users", response_model=list[StaffAccountRead])
def list_users(
    trang_thai: str | None = Query(default=None),
    db: Session = Depends(get_db),
    _admin: User = Depends(require_staff_admin),
):
    selected = (trang_thai or "active").strip().lower()
    query = db.query(User)
    if selected in ("all", "tat ca", "tất cả"):
        pass
    elif selected in ("locked", "nghỉ việc", "nghi viec"):
        query = query.filter(User.trang_thai == "Nghỉ việc")
    elif selected in ("active", "đang làm", "dang lam", ""):
        query = query.filter(or_(User.trang_thai.is_(None), User.trang_thai != "Nghỉ việc"))
    else:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Trạng thái không hợp lệ.")
    try:
        rows = query.order_by(
            case((User.trang_thai == "Nghỉ việc", 1), else_=0).asc(),
            User.id.asc(),
        ).all()
    except SQLAlchemyError as exc:
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from exc
    return [staff_account_read(row) for row in rows]


@app.get("/api/users/roles", response_model=list[str])
def list_user_roles(
    db: Session = Depends(get_db),
    _admin: User = Depends(require_staff_admin),
):
    rows = db.query(User.chuc_vu).distinct().filter(User.chuc_vu.isnot(None)).all()
    titles = []
    seen = set()
    for row in rows:
        title = (row[0] or "").strip()
        key = title.lower()
        if not title or key in seen:
            continue
        seen.add(key)
        titles.append(title)
    titles.sort(key=str.lower)
    return titles


@app.get("/api/users/teams", response_model=list[str])
def list_user_teams(
    db: Session = Depends(get_db),
    _admin: User = Depends(require_staff_admin),
):
    rows = db.query(User.nhom).distinct().filter(User.nhom.isnot(None), User.nhom != "").all()
    teams = []
    seen = set()
    for row in rows:
        name = (row[0] or "").strip()
        key = name.lower()
        if not name or key in seen:
            continue
        seen.add(key)
        teams.append(name)
    teams.sort(key=str.lower)
    return teams


@app.post("/api/users", response_model=StaffAccountRead, status_code=status.HTTP_201_CREATED)
def create_user(
    body: StaffAccountCreate,
    db: Session = Depends(get_db),
    _admin: User = Depends(require_staff_admin),
):
    email = clean_staff_email(body.email)
    exists = db.query(User).filter(func.lower(User.email) == email).first()
    if exists is not None:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Email này đã có trong hệ thống.")
    user = User(
        email=email,
        password_hash=hash_password(body.password),
        full_name=body.full_name.strip(),
        nhom=clean_nhom(body.nhom, db),
    )
    user.role = clean_staff_role(body.role)
    user.chuc_vu = clean_chuc_vu(body.chuc_vu, db)
    user.trang_thai = clean_trang_thai(body.trang_thai)
    db.add(user)
    try:
        db.commit()
        db.refresh(user)
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Email này đã có trong hệ thống.") from exc
    except SQLAlchemyError as exc:
        db.rollback()
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không lưu được nhân viên.",
        ) from exc
    return staff_account_read(user)


@app.put("/api/users/{user_id}", response_model=StaffAccountRead)
def update_user(
    user_id: int,
    body: StaffAccountUpdate,
    db: Session = Depends(get_db),
    _admin: User = Depends(require_staff_admin),
):
    user = find_staff_user(db, user_id)
    email = clean_staff_email(body.email)
    role = clean_staff_role(body.role)
    taken = (
        db.query(User)
        .filter(func.lower(User.email) == email, User.id != user.id)
        .first()
    )
    if taken is not None:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Email này đã có trong hệ thống.")
    if (user.role or "").strip().lower() == "admin" and role != "admin" and admin_count(db, exclude_id=user.id) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không thể hạ quyền quản trị viên cuối cùng.",
        )
    user.full_name = body.full_name.strip()
    user.email = email
    user.role = role
    user.chuc_vu = clean_chuc_vu(body.chuc_vu, db)
    user.nhom = clean_nhom(body.nhom, db)
    next_status = clean_trang_thai(body.trang_thai)
    if next_status == "Nghỉ việc":
        if user.id == _admin.id:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Không thể khóa tài khoản đang đăng nhập.")
        if (user.role or "").strip().lower() == "admin" and admin_count(db, exclude_id=user.id) == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Không thể khóa quản trị viên cuối cùng.",
            )
    user.trang_thai = next_status
    try:
        db.commit()
        db.refresh(user)
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Email này đã có trong hệ thống.") from exc
    except SQLAlchemyError as exc:
        db.rollback()
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không cập nhật được nhân viên.",
        ) from exc
    return staff_account_read(user)


@app.delete("/api/users/{user_id}", response_model=MessageResponse)
def delete_user(
    user_id: int,
    db: Session = Depends(get_db),
    admin: User = Depends(require_staff_admin),
):
    user = find_staff_user(db, user_id)
    if user.id == admin.id:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Không thể khóa tài khoản đang đăng nhập.")
    if (user.role or "").strip().lower() == "admin" and admin_count(db, exclude_id=user.id) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không thể khóa quản trị viên cuối cùng.",
        )
    user.trang_thai = "Nghỉ việc"
    try:
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không khóa được tài khoản.",
        ) from exc
    return MessageResponse(message="Đã khóa tài khoản.")


@app.post("/api/users/{user_id}/reset-password", response_model=ResetPasswordResponse)
def reset_user_password(
    user_id: int,
    db: Session = Depends(get_db),
    _admin: User = Depends(require_staff_admin),
):
    user = find_staff_user(db, user_id)
    user.password_hash = hash_password(DEFAULT_STAFF_PASSWORD)
    try:
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không đặt lại được mật khẩu.",
        ) from exc
    return ResetPasswordResponse(
        message="Đã đặt lại mật khẩu.",
        default_password=DEFAULT_STAFF_PASSWORD,
    )


@app.get("/api/properties/{property_id}/comments", response_model=list[CommentRead])
def list_comments(
    property_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
):
    found = db.get(Property, property_id)
    if found is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy bất động sản.")
    rows = (
        db.query(Comment)
        .options(joinedload(Comment.author))
        .filter(Comment.property_id == property_id)
        .order_by(Comment.created_at.asc(), Comment.id.asc())
        .all()
    )
    return [comment_read(row) for row in rows]


@app.post("/api/properties/{property_id}/comments", response_model=CommentRead, status_code=status.HTTP_201_CREATED)
def create_comment(
    property_id: int,
    body: CommentCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    found = db.get(Property, property_id)
    if found is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy bất động sản.")
    item = Comment(
        property_id=property_id,
        user_id=current_user.id,
        content=body.content,
        created_at=datetime.now(),
    )
    db.add(item)
    try:
        db.commit()
        db.refresh(item)
    except SQLAlchemyError as exc:
        db.rollback()
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không lưu được bình luận.",
        ) from exc
    item.author = current_user
    return comment_read(item)


def require_property_owner(item: Property, user: User) -> None:
    if item.user_id is not None and item.user_id == user.id:
        return
    if (user.role or "").strip().lower() == "admin":
        return
    raise HTTPException(
        status_code=status.HTTP_403_FORBIDDEN,
        detail="Bạn không có quyền sửa/xóa bài viết của người khác",
    )


def author_label(item: Property) -> str:
    author = item.author
    if author is None:
        return "Nhân viên"
    return author.full_name or author.email or "Nhân viên"


def property_title(content: str | None) -> str:
    text = content or ""
    match = re.search(r"(?im)^\s*1\.\s*(?:Tiêu đề\s*:\s*)?(.*?)\s*$", text)
    if match:
        title = re.sub(r"(?i)^tiêu đề\s*:\s*", "", match.group(1)).strip()
        if title:
            return title[:255]
    for line in text.splitlines():
        cleaned = line.strip()
        if cleaned:
            return cleaned[:255]
    return "Bài viết"


def find_parcel_match(db: Session, so_to: str | None, so_thua: str | None, exclude_id: int | None = None) -> Property | None:
    sheet = (so_to or "").strip()
    plot = (so_thua or "").strip()
    if not sheet or not plot:
        return None
    query = (
        db.query(Property)
        .options(joinedload(Property.author))
        .filter(Property.so_to == sheet, Property.so_thua == plot)
    )
    if exclude_id is not None:
        query = query.filter(Property.id != exclude_id)
    return query.first()


def duplicate_message(item: Property, sheet: str, plot: str) -> str:
    return (
        "⚠️ Bất động sản này (Số tờ "
        + sheet
        + ", Số thửa "
        + plot
        + ") đã được đăng bởi nhân viên "
        + author_label(item)
        + ". Không thể tạo mới!"
    )


def to_property_read(item: Property) -> PropertyRead:
    author = item.author
    name = ""
    if author is not None:
        name = author.full_name or author.email
    return PropertyRead(
        id=item.id,
        content=item.content or "",
        price=float(item.price) if item.price is not None else None,
        area=float(item.area) if item.area is not None else None,
        location=item.location or "",
        so_to=item.so_to,
        so_thua=item.so_thua,
        direction=item.direction,
        media_urls=read_media_urls(item.media_urls),
        video_url=item.video_url or None,
        status=item.status,
        created_at=item.created_at,
        user_id=item.user_id,
        author_name=name,
    )


@app.get("/api/properties", response_model=list[PropertyRead])
def list_properties(
    user_id: int | None = Query(default=None),
    min_price: float | None = Query(default=None),
    max_price: float | None = Query(default=None),
    direction: str | None = Query(default=None),
    address: str | None = Query(default=None),
    location: str | None = Query(default=None),
    so_thua: str | None = Query(default=None),
    so_to: str | None = Query(default=None),
    status_filter: str | None = Query(default=None, alias="status"),
    author: str | None = Query(default=None),
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    try:
        query = db.query(Property).options(joinedload(Property.author))
        if user_id is not None:
            query = query.filter(Property.user_id == user_id)
        if min_price is not None:
            query = query.filter(Property.price >= min_price)
        if max_price is not None:
            query = query.filter(Property.price <= max_price)
        if direction:
            query = query.filter(Property.direction == direction.strip())
        place = (address or location or "").strip()
        if place:
            query = query.filter(Property.location.ilike(f"%{place}%"))
        plot = (so_thua or "").strip()
        if plot:
            query = query.filter(Property.so_thua.ilike(f"%{plot}%"))
        sheet = (so_to or "").strip()
        if sheet:
            query = query.filter(Property.so_to.ilike(f"%{sheet}%"))
        sale_status = (status_filter or "").strip()
        if sale_status:
            query = query.filter(Property.status == sale_status)
        poster = (author or "").strip()
        if poster:
            query = query.filter(Property.author.has(User.full_name.ilike(f"%{poster}%")))
        rows = query.order_by(Property.created_at.desc()).all()
    except SQLAlchemyError as e:
        print(f"Chi tiết lỗi DB: {e}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from e
    return [to_property_read(item) for item in rows]


@app.get("/api/properties/check-duplicate", response_model=DuplicateMatch | None)
def check_duplicate_property(
    so_to: str = Query(default=""),
    so_thua: str = Query(default=""),
    db: Session = Depends(get_db),
    _user: User = Depends(get_current_user),
):
    sheet = (so_to or "").strip()
    plot = (so_thua or "").strip()
    if not sheet or not plot:
        return None
    try:
        item = find_parcel_match(db, sheet, plot)
    except SQLAlchemyError as exc:
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from exc
    if item is None:
        return None
    return DuplicateMatch(
        id=item.id,
        title=property_title(item.content),
        author_name=author_label(item),
    )


def pending_image_path(value: str) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    path = urllib.parse.urlparse(text).path if "://" in text or text.startswith("//") else text
    marker = path.find("/static/")
    if marker < 0:
        marker = text.find("/static/")
        path = text[marker:] if marker >= 0 else ""
    else:
        path = path[marker:]
    path = path.split("?", 1)[0].split("#", 1)[0]
    if not path.startswith("/static/"):
        return ""
    return path[:1000]


class PendingPostCreate(BaseModel):
    content: str = Field(min_length=1, max_length=20000)
    images: list[str] = Field(default_factory=list)
    video_url: str = ""
    employee_name: str = ""
    status: str = "pending"

    @field_validator("content")
    @classmethod
    def clean_pending_content(cls, value: str) -> str:
        text = (value or "").strip()
        if not text:
            raise ValueError("Nội dung bài chờ duyệt không được để trống.")
        return text

    @field_validator("images")
    @classmethod
    def clean_pending_images(cls, value: list[str]) -> list[str]:
        cleaned: list[str] = []
        for item in value or []:
            path = pending_image_path(str(item or ""))
            if path and path not in cleaned:
                cleaned.append(path)
        return cleaned[:40]

    @field_validator("video_url")
    @classmethod
    def clean_pending_video(cls, value: str) -> str:
        return pending_image_path(str(value or ""))

    @field_validator("employee_name")
    @classmethod
    def clean_employee_name(cls, value: str) -> str:
        return (value or "").strip()[:255]

    @field_validator("status")
    @classmethod
    def force_pending_status(cls, _value: str) -> str:
        return "pending"


class PendingPostRead(BaseModel):
    id: int
    content: str
    images: list[str]
    video_url: str | None = None
    employee_name: str
    status: str
    user_id: int | None
    created_at: datetime


@app.post("/api/pending-posts", response_model=PendingPostRead, status_code=status.HTTP_201_CREATED)
def create_pending_post(
    body: PendingPostCreate,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    employee_name = (current_user.full_name or "").strip() or body.employee_name or "Nhân viên"
    item = PendingPost(
        user_id=current_user.id,
        employee_name=employee_name[:255],
        content=body.content,
        images=json.dumps(body.images, ensure_ascii=False),
        video_url=body.video_url or None,
        status="pending",
        created_at=datetime.now(),
    )
    try:
        db.add(item)
        db.commit()
        db.refresh(item)
    except SQLAlchemyError as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không lưu được bài chờ duyệt.",
        ) from exc
    return pending_post_read(item)


def pending_post_read(item: PendingPost) -> PendingPostRead:
    return PendingPostRead(
        id=item.id,
        content=item.content,
        images=read_media_urls(item.images),
        video_url=item.video_url or None,
        employee_name=item.employee_name,
        status=item.status,
        user_id=item.user_id,
        created_at=item.created_at,
    )


def require_admin(user: User) -> None:
    if not is_system_admin(user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Bạn không có quyền truy cập")


def find_pending_post(db: Session, post_id: int) -> PendingPost:
    item = db.get(PendingPost, post_id)
    if item is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy bài chờ duyệt.")
    if item.status != "pending":
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Bài viết không còn ở trạng thái chờ duyệt.")
    return item


class PendingPostsCount(BaseModel):
    count: int


@app.get("/api/dashboard/pending-posts-count", response_model=PendingPostsCount)
def pending_posts_count(
    db: Session = Depends(get_db),
    _user: User = Depends(require_module_admin),
):
    total = (
        db.query(func.count(PendingPost.id))
        .filter(PendingPost.status == "pending")
        .scalar()
        or 0
    )
    return PendingPostsCount(count=int(total))


@app.get("/api/pending-posts", response_model=List[PendingPostRead])
def list_pending_posts(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_admin(current_user)
    rows = (
        db.query(PendingPost)
        .filter(PendingPost.status == "pending")
        .order_by(PendingPost.created_at.desc(), PendingPost.id.desc())
        .all()
    )
    return [pending_post_read(row) for row in rows]


class PendingApproveRequest(BaseModel):
    images: list[str] = Field(default_factory=list)
    media_choice: str = "image"

    @field_validator("images")
    @classmethod
    def clean_approved_images(cls, value: list[str]) -> list[str]:
        cleaned: list[str] = []
        for item in value or []:
            path = pending_image_path(str(item or ""))
            if path and path not in cleaned:
                cleaned.append(path)
        return cleaned[:40]

    @field_validator("media_choice")
    @classmethod
    def clean_media_choice(cls, value: str) -> str:
        choice = (value or "image").strip().lower()
        if choice not in {"image", "video"}:
            return "image"
        return choice


class FanpagePublishResponse(BaseModel):
    message: str


@app.post("/api/approve-post/{post_id}", response_model=FanpagePublishResponse)
def approve_post(
    post_id: int,
    body: PendingApproveRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_admin(current_user)
    load_dotenv(Path(__file__).resolve().parent / ".env", override=True)
    webhook_url = (os.getenv("MAKE_WEBHOOK_URL") or "").strip().strip('"').strip("'")
    if not webhook_url:
        print("Debug: MAKE_WEBHOOK_URL đang bị None")
        raise HTTPException(status_code=400, detail="Chưa cấu hình webhook Make.com")
    item = find_pending_post(db, post_id)
    stored_raw = read_media_urls(item.images)
    stored = [pending_image_path(url) for url in stored_raw]
    allowed = {url for url in stored if url}
    kept = [url for url in body.images if url in allowed]
    original_by_path = {}
    for raw in stored_raw:
        path = pending_image_path(raw)
        if path:
            original_by_path[path] = raw
    public_domain = "crm.nhadatthanhtam.com"
    images = []
    for path in kept:
        original_url = original_by_path.get(path, path)
        if "localhost" in original_url:
            cleaned_url = original_url.replace("http://localhost:8000", public_domain)
        elif original_url.startswith("/"):
            cleaned_url = public_domain + original_url
        else:
            cleaned_url = original_url
        images.append(cleaned_url)
    raw_video = (item.video_url or "").strip()
    if "localhost" in raw_video:
        public_video_url = raw_video.replace("http://localhost:8000", public_domain)
    elif raw_video.startswith("/"):
        public_video_url = public_domain + raw_video
    else:
        public_video_url = raw_video
    if body.media_choice == "video":
        images = []
    else:
        public_video_url = ""
    payload = {
        "post_id": item.id,
        "content": item.content or "",
        "images": images,
        "video_url": public_video_url,
    }
    try:
        webhook_response = requests.post(webhook_url, json=payload, timeout=45)
    except requests.RequestException as exc:
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail="Không gửi được bài sang Fanpage.",
        ) from exc
    if webhook_response.status_code != 200:
        raise HTTPException(status_code=status.HTTP_502_BAD_GATEWAY, detail="Fanpage chưa nhận được bài.")
    item.images = json.dumps(kept, ensure_ascii=False)
    item.status = "published"
    try:
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không lưu được trạng thái đã đăng.",
        ) from exc
    return FanpagePublishResponse(message="Đăng Fanpage thành công!")


@app.delete("/api/pending-posts/{post_id}")
def delete_pending_post(
    post_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_admin(current_user)
    item = db.get(PendingPost, post_id)
    if item is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy bài chờ duyệt.")
    try:
        db.delete(item)
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không từ chối được bài viết.",
        ) from exc
    return {"message": "Đã từ chối và xóa bài viết"}


@app.post("/api/pending-posts/{post_id}/reject", response_model=PendingPostRead)
def reject_pending_post(
    post_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_admin(current_user)
    item = find_pending_post(db, post_id)
    item.status = "rejected"
    try:
        db.commit()
        db.refresh(item)
    except SQLAlchemyError as exc:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không từ chối được bài viết.",
        ) from exc
    return pending_post_read(item)


@app.post("/api/generate-content", response_model=ComposeResponse)
def generate_content(body: ComposeRequest, _user: User = Depends(get_current_user)):
    return ComposeResponse(content=rewrite_with_openai(body.content, body.style))


@app.post("/api/properties/compose", response_model=ComposeResponse)
def compose_property(body: ComposeRequest, _user: User = Depends(get_current_user)):
    return ComposeResponse(content=rewrite_with_openai(body.content, body.style))


@app.post("/api/properties", response_model=PropertyRead, status_code=status.HTTP_201_CREATED)
def create_property(
    content: str = Form(...),
    price: str = Form(""),
    area: str = Form(""),
    location: str = Form(""),
    so_to: str = Form(""),
    so_thua: str = Form(""),
    direction: str = Form(""),
    status_value: str | None = Form(default=None),
    media_files: List[UploadFile] | None = File(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    text_body = content.strip()
    if not text_body:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Hãy dán nội dung bài viết.",
        )
    price_value = optional_float(price, "Giá")
    area_value = optional_float(area, "Diện tích")
    create_data = {
        "content": text_body,
        "price": price_value,
        "area": area_value,
        "location": location.strip()[:255],
        "so_to": so_to,
        "so_thua": so_thua,
        "direction": normalize_direction(direction) or parse_direction(text_body),
    }
    if status_value is not None and status_value.strip():
        create_data["status"] = status_value.strip()
    try:
        created = PropertyCreate(**create_data)
    except ValidationError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=exc.errors()[0]["msg"] if exc.errors() else "Dữ liệu không hợp lệ.",
        ) from exc
    try:
        existing = find_parcel_match(db, created.so_to, created.so_thua)
    except SQLAlchemyError as exc:
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from exc
    if existing is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail=duplicate_message(existing, created.so_to or "", created.so_thua or ""),
        )
    stored: list[Path] = []
    video_stored: list[Path] = []
    try:
        image_uploads, video_upload = classify_media_uploads(media_files)
        urls, stored = save_media_files(image_uploads)
        video_path, video_stored = save_video_file(video_upload)
    except HTTPException:
        for path in stored + video_stored:
            path.unlink(missing_ok=True)
        raise
    except OSError as exc:
        print(f"Chi tiết lỗi lưu file: {exc}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Không lưu được ảnh hoặc video.",
        ) from exc
    item = Property(
        user_id=current_user.id,
        content=created.content,
        price=created.price,
        area=created.area,
        location=created.location or None,
        so_to=created.so_to,
        so_thua=created.so_thua,
        direction=created.direction,
        media_urls=json.dumps(urls, ensure_ascii=False),
        video_url=video_path or None,
        status=created.status,
        created_at=datetime.now(),
    )
    try:
        db.add(item)
        db.commit()
        db.refresh(item)
        item.author = current_user
    except SQLAlchemyError as e:
        db.rollback()
        for path in stored + video_stored:
            path.unlink(missing_ok=True)
        print(f"Chi tiết lỗi DB: {e}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from e
    return to_property_read(item)


@app.put("/api/properties/{property_id}", response_model=PropertyRead)
def update_property(
    property_id: int,
    content: str = Form(...),
    status_value: str = Form(..., alias="status"),
    price: str = Form(""),
    area: str = Form(""),
    location: str = Form(""),
    direction: str = Form(""),
    so_to: str | None = Form(default=None),
    so_thua: str | None = Form(default=None),
    keep_media: List[str] | None = Form(default=None),
    media_files: List[UploadFile] | None = File(None),
    clear_video: str = Form(""),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    text_body = content.strip()
    if not text_body:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Hãy nhập nội dung bài viết.",
        )
    update_data = {
        "content": text_body,
        "price": optional_float(price, "Giá"),
        "area": optional_float(area, "Diện tích"),
        "location": location.strip()[:255],
        "status": status_value,
        "direction": normalize_direction(direction) or parse_direction(text_body),
    }
    if so_to is not None:
        update_data["so_to"] = so_to
    if so_thua is not None:
        update_data["so_thua"] = so_thua
    try:
        updated = PropertyCreate(**update_data)
    except ValidationError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=exc.errors()[0]["msg"] if exc.errors() else "Dữ liệu không hợp lệ.",
        ) from exc
    try:
        item = (
            db.query(Property)
            .options(joinedload(Property.author))
            .filter(Property.id == property_id)
            .first()
        )
    except SQLAlchemyError as e:
        print(f"Chi tiết lỗi DB: {e}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from e
    if item is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy bài viết.")
    require_property_owner(item, current_user)
    next_sheet = updated.so_to if so_to is not None else item.so_to
    next_plot = updated.so_thua if so_thua is not None else item.so_thua
    if so_to is not None or so_thua is not None:
        try:
            existing = find_parcel_match(db, next_sheet, next_plot, exclude_id=item.id)
        except SQLAlchemyError as exc:
            print(f"Chi tiết lỗi DB: {exc}")
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="Không kết nối được cơ sở dữ liệu.",
            ) from exc
        if existing is not None:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=duplicate_message(existing, next_sheet or "", next_plot or ""),
            )
    old_urls = read_media_urls(item.media_urls)
    requested = keep_media or []
    kept = [url for url in old_urls if url in requested]
    removed = [url for url in old_urls if url not in kept]
    stored: list[Path] = []
    video_stored: list[Path] = []
    try:
        image_uploads, video_upload = classify_media_uploads(media_files)
        new_urls, stored = save_media_files(image_uploads)
        video_path, video_stored = save_video_file(video_upload)
    except HTTPException:
        for path in stored + video_stored:
            path.unlink(missing_ok=True)
        raise
    except OSError as exc:
        print(f"Chi tiết lỗi lưu file: {exc}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Không lưu được ảnh hoặc video.",
        ) from exc
    final_urls = kept + new_urls
    if len(final_urls) > 12:
        for path in stored + video_stored:
            path.unlink(missing_ok=True)
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Mỗi tin tối đa 12 ảnh hoặc video.",
        )
    removed_video = ""
    if video_path:
        if item.video_url and item.video_url != video_path:
            removed_video = item.video_url
        item.video_url = video_path
    elif (clear_video or "").strip() in {"1", "true", "yes"}:
        if item.video_url:
            removed_video = item.video_url
        item.video_url = None
    item.content = updated.content
    item.status = updated.status
    item.price = updated.price
    item.area = updated.area
    item.location = updated.location or None
    item.direction = updated.direction
    if so_to is not None:
        item.so_to = updated.so_to
    if so_thua is not None:
        item.so_thua = updated.so_thua
    item.media_urls = json.dumps(final_urls, ensure_ascii=False)
    author = item.author
    try:
        db.commit()
        db.refresh(item)
        item.author = author
    except SQLAlchemyError as e:
        db.rollback()
        for path in stored + video_stored:
            path.unlink(missing_ok=True)
        print(f"Chi tiết lỗi DB: {e}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from e
    for url in removed:
        path = local_upload(url)
        if path is not None:
            path.unlink(missing_ok=True)
    if removed_video:
        old_video_path = local_upload(removed_video)
        if old_video_path is not None:
            old_video_path.unlink(missing_ok=True)
    return to_property_read(item)


@app.delete("/api/properties/{property_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_property(
    property_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    try:
        item = db.query(Property).filter(Property.id == property_id).first()
    except SQLAlchemyError as exc:
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không kết nối được cơ sở dữ liệu.",
        ) from exc
    if item is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy bài viết.")
    require_property_owner(item, current_user)
    urls = read_media_urls(item.media_urls)
    video_url = item.video_url or ""
    try:
        db.delete(item)
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        print(f"Chi tiết lỗi DB: {exc}")
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Không xóa được bài viết.",
        ) from exc
    for url in urls:
        path = local_upload(url)
        if path is not None:
            path.unlink(missing_ok=True)
    video_path = local_upload(video_url)
    if video_path is not None:
        video_path.unlink(missing_ok=True)
    return None


def frontend_page(name: str) -> FileResponse:
    root = FRONTEND_ROOT.resolve()
    path = (root / name).resolve()
    if path.parent != root or not path.is_file():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Không tìm thấy trang.")
    return FileResponse(path)


@app.get("/", include_in_schema=False)
def frontend_home():
    return frontend_page("index.html")


@app.get("/index.html", include_in_schema=False)
def frontend_index():
    return frontend_page("index.html")


@app.get("/pending-posts.html", include_in_schema=False)
def frontend_pending_posts():
    return frontend_page("pending-posts.html")


@app.get("/properties.html", include_in_schema=False)
def frontend_properties():
    return frontend_page("properties.html")


@app.get("/dashboard.html", include_in_schema=False)
def frontend_dashboard():
    return frontend_page("dashboard.html")


@app.get("/users.html", include_in_schema=False)
def frontend_users():
    return frontend_page("users.html")


@app.get("/land-prices.html", include_in_schema=False)
def frontend_land_prices():
    return frontend_page("land-prices.html")


@app.get("/contracts.html", include_in_schema=False)
def frontend_contracts():
    return frontend_page("contracts.html")


@app.get("/tracuu")
def trang_tra_cuu():
    # Lấy đường dẫn chính xác tới thư mục frontend (do backend và frontend nằm ngang hàng trong QuanLyBDS)
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    html_path = os.path.join(base_dir, "frontend", "tracuu.html")

    # Trả về file giao diện cho trình duyệt
    return FileResponse(html_path)


app.mount("/static", StaticFiles(directory=str(STATIC_ROOT)), name="static")
app.mount("/assets", StaticFiles(directory=str(FRONTEND_ROOT / "assets")), name="assets")
